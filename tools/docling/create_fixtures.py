from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def make_pdf(path: Path) -> None:
    stream = b"BT /F1 18 Tf 72 720 Td (Boss Office PDF document parsing test) Tj 0 -30 Td /F1 12 Tf (Local Docling extraction works without cloud services.) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects)+1}\n".encode())
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


def main() -> None:
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/verification/docling-tests/fixtures")
    target.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("老闆辦公室文件解析測試", 0)
    document.add_paragraph("這份 DOCX 用來確認 Docling 能在本機讀取繁體中文，並把內容交給主管。")
    document.add_heading("驗收條件", 1)
    document.add_paragraph("不得連接雲端，不得啟動 Codex 模型，關閉頁面時必須停止解析。")
    document.save(target / "phase2-sample.docx")
    make_pdf(target / "phase2-sample.pdf")


if __name__ == "__main__":
    main()
